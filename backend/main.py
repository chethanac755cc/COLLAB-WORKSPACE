import os
import uuid
import io
import csv

from sentence_transformers import SentenceTransformer
import httpx
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, WebSocket, WebSocketDisconnect, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from supabase import create_client, Client
from fastapi.responses import Response

load_dotenv()

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_SECRET_KEY = os.getenv("SUPABASE_SECRET_KEY")

supabase: Client = create_client(
    SUPABASE_URL,
    SUPABASE_SECRET_KEY
)
embedding_model = SentenceTransformer("all-MiniLM-L6-v2")



app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


connections = {}


@app.get("/")
def home():
    return {
        "message": "Collaborative Workspace API is running!"
    }


@app.get("/api/test")
def test():
    return {
        "message": "Hello from FastAPI!"
    }


@app.post("/api/rooms")
def create_room():
    room_id = uuid.uuid4().hex[:6].upper()

    # Create room
    room_response = (
        supabase
        .table("rooms")
        .insert({
            "room_code": room_id
        })
        .execute()
    )

    room_db_id = room_response.data[0]["id"]

    # Create default code file
    supabase.table("code_documents").insert({
        "room_id": room_db_id,
        "filename": "main.js",
        "code": "// Start coding together..."
    }).execute()

    return {
        "room_id": room_id
    }

@app.get("/api/rooms/{room_id}")
def join_room(room_id: str):
    room_id = room_id.upper()

    response = (
        supabase
        .table("rooms")
        .select("*")
        .eq("room_code", room_id)
        .execute()
    )

    if not response.data:
        raise HTTPException(
            status_code=404,
            detail="Room not found"
        )

    return {
    "room_id": response.data[0]["room_code"]
}

@app.get("/api/rooms/{room_id}/code")
def get_code(room_id: str):
    room_id = room_id.upper()

    try:
        room_response = (
            supabase
            .table("rooms")
            .select("id")
            .eq("room_code", room_id)
            .execute()
        )

        if not room_response.data:
            raise HTTPException(
                status_code=404,
                detail="Room not found"
            )

        room_db_id = room_response.data[0]["id"]

        code_response = (
            supabase
            .table("code_documents")
            .select("*")
            .eq("room_id", room_db_id)
            .limit(1)
            .execute()
        )

        if not code_response.data:
            raise HTTPException(
                status_code=404,
                detail="Code file not found"
            )

        return code_response.data[0]

    except HTTPException:
        raise

    except Exception as error:
        print("Database error while loading code:", error)

        raise HTTPException(
            status_code=503,
            detail="Database temporarily unavailable"
        )


from pydantic import BaseModel

class CodeUpdate(BaseModel):
    code: str


@app.put("/api/rooms/{room_id}/code")
def update_code(room_id: str, data: CodeUpdate):
    room_id = room_id.upper()

    try:
        # Find the room
        room_response = (
            supabase
            .table("rooms")
            .select("id")
            .eq("room_code", room_id)
            .execute()
        )

        if not room_response.data:
            raise HTTPException(
                status_code=404,
                detail="Room not found"
            )

        room_db_id = room_response.data[0]["id"]

        # Find the code document
        code_response = (
            supabase
            .table("code_documents")
            .select("id")
            .eq("room_id", room_db_id)
            .limit(1)
            .execute()
        )

        if not code_response.data:
            raise HTTPException(
                status_code=404,
                detail="Code file not found"
            )

        code_id = code_response.data[0]["id"]

        # Update code
        supabase \
            .table("code_documents") \
            .update({
                "code": data.code,
                "updated_at": "now()"
            }) \
            .eq("id", code_id) \
            .execute()

        return {
            "saved": True
        }

    except HTTPException:
        raise

    except Exception as error:
        print("Error saving code:", error)

        raise HTTPException(
            status_code=503,
            detail="Could not save code"
        )

@app.websocket("/ws/{room_id}")
async def websocket_endpoint(websocket: WebSocket, room_id: str):
    room_id = room_id.upper()

    await websocket.accept()

    connections.setdefault(room_id, []).append(websocket)

    try:
        while True:
            message = await websocket.receive_text()

            for connection in connections.get(room_id, []):
                if connection != websocket:
                    await connection.send_text(message)

    except WebSocketDisconnect:
        if websocket in connections.get(room_id, []):
            connections[room_id].remove(websocket)

        if not connections.get(room_id):
            connections.pop(room_id, None)


@app.get("/api/db-test")
def db_test():
    response = supabase.table("rooms").select("*").limit(1).execute()

    return {
        "connected": True,
        "data": response.data
    }

@app.get("/api/rooms/{room_id}/files")
def get_files(room_id: str):
    room_id = room_id.upper()

    try:
        room_response = (
            supabase
            .table("rooms")
            .select("id")
            .eq("room_code", room_id)
            .execute()
        )

        if not room_response.data:
            raise HTTPException(
                status_code=404,
                detail="Room not found"
            )

        room_db_id = room_response.data[0]["id"]

        response = (
            supabase
            .table("code_documents")
            .select("id, filename, updated_at")
            .eq("room_id", room_db_id)
            .order("filename")
            .execute()
        )

        return response.data

    except HTTPException:
        raise

    except Exception as error:
        print("Error loading files:", error)

        raise HTTPException(
            status_code=503,
            detail="Could not load files"
        )


@app.post("/api/rooms/{room_id}/files")
def create_file(room_id: str, data: dict):
    room_id = room_id.upper()

    filename = data.get("filename", "").strip()

    if not filename:
        raise HTTPException(
            status_code=400,
            detail="Filename is required"
        )

    try:
        room_response = (
            supabase
            .table("rooms")
            .select("id")
            .eq("room_code", room_id)
            .execute()
        )

        if not room_response.data:
            raise HTTPException(
                status_code=404,
                detail="Room not found"
            )

        room_db_id = room_response.data[0]["id"]

        existing = (
            supabase
            .table("code_documents")
            .select("id")
            .eq("room_id", room_db_id)
            .eq("filename", filename)
            .execute()
        )

        if existing.data:
            raise HTTPException(
                status_code=400,
                detail="File already exists"
            )

        response = (
            supabase
            .table("code_documents")
            .insert({
                "room_id": room_db_id,
                "filename": filename,
                "code": ""
            })
            .execute()
        )

        return response.data[0]

    except HTTPException:
        raise

    except Exception as error:
        print("Error creating file:", error)

        raise HTTPException(
            status_code=503,
            detail="Could not create file"
        )


@app.get("/api/files/{file_id}")
def get_file(file_id: str):
    try:
        response = (
            supabase
            .table("code_documents")
            .select("*")
            .eq("id", file_id)
            .execute()
        )

        if not response.data:
            raise HTTPException(
                status_code=404,
                detail="File not found"
            )

        return response.data[0]

    except HTTPException:
        raise

    except Exception as error:
        print("Error loading file:", error)

        raise HTTPException(
            status_code=503,
            detail="Could not load file"
        )


@app.put("/api/files/{file_id}")
def update_file(file_id: str, data: CodeUpdate):
    try:
        response = (
            supabase
            .table("code_documents")
            .update({
                "code": data.code,
                "updated_at": "now()"
            })
            .eq("id", file_id)
            .execute()
        )

        if not response.data:
            raise HTTPException(
                status_code=404,
                detail="File not found"
            )

        return {
            "saved": True
        }

    except HTTPException:
        raise

    except Exception as error:
        print("Error saving file:", error)

        raise HTTPException(
            status_code=503,
            detail="Could not save file"
        )


@app.delete("/api/files/{file_id}")
def delete_file(file_id: str):
    try:
        response = (
            supabase
            .table("code_documents")
            .delete()
            .eq("id", file_id)
            .execute()
        )

        return {
            "deleted": True
        }

    except Exception as error:
        print("Error deleting file:", error)

        raise HTTPException(
            status_code=503,
            detail="Could not delete file"
        )

def extract_text(file_data: bytes, filename: str) -> str:
    """
    Extract readable text from common document formats.
    """

    extension = filename.lower().split(".")[-1]

    # -----------------------------
    # PDF
    # -----------------------------
    if extension == "pdf":
        reader = PdfReader(io.BytesIO(file_data))

        pages = []

        for page in reader.pages:
            text = page.extract_text()

            if text:
                pages.append(text)

        return "\n\n".join(pages)

    # -----------------------------
    # DOCX
    # -----------------------------
    if extension == "docx":
        document = Document(io.BytesIO(file_data))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        return "\n".join(paragraphs)

    # -----------------------------
    # XLSX
    # -----------------------------
    if extension == "xlsx":
        workbook = load_workbook(
            io.BytesIO(file_data),
            data_only=True
        )

        lines = []

        for sheet in workbook.worksheets:
            lines.append(f"Sheet: {sheet.title}")

            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value)
                    for value in row
                    if value is not None
                ]

                if values:
                    lines.append(" | ".join(values))

        return "\n".join(lines)

    # -----------------------------
    # PPTX
    # -----------------------------
    if extension == "pptx":
        presentation = Presentation(
            io.BytesIO(file_data)
        )

        slides = []

        for index, slide in enumerate(
            presentation.slides,
            start=1
        ):
            slide_text = []

            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text)

            if slide_text:
                slides.append(
                    f"Slide {index}\n"
                    + "\n".join(slide_text)
                )

        return "\n\n".join(slides)

    # -----------------------------
    # TXT
    # -----------------------------
    if extension == "txt":
        return file_data.decode(
            "utf-8",
            errors="ignore"
        )

    # -----------------------------
    # CSV
    # -----------------------------
    if extension == "csv":
        text = file_data.decode(
            "utf-8",
            errors="ignore"
        )

        reader = csv.reader(
            io.StringIO(text)
        )

        rows = []

        for row in reader:
            rows.append(" | ".join(row))

        return "\n".join(rows)

    # Unsupported format
    return ""

def create_chunks(text: str, chunk_size: int = 1000, overlap: int = 200):
    """
    Split extracted text into overlapping chunks.
    """

    text = text.strip()

    if not text:
        return []

    chunks = []

    start = 0

    while start < len(text):
        end = start + chunk_size

        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        start += chunk_size - overlap

    return chunks


def generate_embedding(text: str):
    return embedding_model.encode(text).tolist()

def search_document_chunks(
    room_db_id: str,
    query: str,
    match_count: int = 5
):
    query_embedding = generate_embedding(query)

    response = (
        supabase.rpc(
            "match_document_chunks",
            {
                "query_embedding": query_embedding,
                "room_uuid": room_db_id,
                "match_threshold": 0.20,
                "match_count": match_count
            }
        )
        .execute()
    )

    return response.data or []

  
@app.post("/api/rooms/{room_id}/upload")
async def upload_file(
    room_id: str,
    uploaded_file: UploadFile = File(...)
):
    room_id = room_id.upper()

    try:
        # -----------------------------------------
        # FIND ROOM
        # -----------------------------------------

        room_response = (
            supabase
            .table("rooms")
            .select("id")
            .eq("room_code", room_id)
            .execute()
        )

        if not room_response.data:
            raise HTTPException(
                status_code=404,
                detail="Room not found"
            )

        room_db_id = room_response.data[0]["id"]

        # -----------------------------------------
        # GET FILENAME
        # -----------------------------------------

        filename = uploaded_file.filename

        if not filename:
            raise HTTPException(
                status_code=400,
                detail="Filename is required"
            )

        # -----------------------------------------
        # READ FILE
        # -----------------------------------------

        file_data = await uploaded_file.read()

        print("========== EXTRACTION START ==========")
        print("Filename:", filename)
        print("File size:", len(file_data))

        # -----------------------------------------
        # EXTRACT TEXT
        # -----------------------------------------

        extracted_text = extract_text(
            file_data,
            filename
        )

        print(
            "Extracted characters:",
            len(extracted_text)
        )

        print(
            "Preview:",
            extracted_text[:500]
        )

        print("========== EXTRACTION END ==========")

        # -----------------------------------------
        # CREATE UNIQUE STORAGE PATH
        # -----------------------------------------

        storage_path = (
            f"{room_id}/{uuid.uuid4()}_{filename}"
        )

        # -----------------------------------------
        # UPLOAD TO SUPABASE STORAGE
        # -----------------------------------------

        supabase.storage \
            .from_("project-files") \
            .upload(
                storage_path,
                file_data,
                {
                    "content-type":
                        uploaded_file.content_type
                        or "application/octet-stream"
                }
            )

        # -----------------------------------------
        # GET PUBLIC URL
        # -----------------------------------------

        file_url = (
            supabase.storage
            .from_("project-files")
            .get_public_url(storage_path)
        )

        # -----------------------------------------
        # SAVE FILE METADATA
        # -----------------------------------------

        response = (
            supabase
            .table("files")
            .insert({
                "room_id": room_db_id,
                "filename": filename,
                "file_url": file_url,
                "file_type":
                    uploaded_file.content_type
                    or "application/octet-stream"
            })
            .execute()
        )

        file_record = response.data[0]

        # Get database file ID
        file_id = file_record["id"]

        # -----------------------------------------
        # CREATE TEXT CHUNKS
        # -----------------------------------------

        chunks = create_chunks(
            extracted_text
        )

        print(
            f"Created {len(chunks)} chunks from {filename}"
        )

        # -----------------------------------------
        # GENERATE EMBEDDINGS
        # -----------------------------------------

        if chunks:

            chunk_records = []

            for index, chunk in enumerate(chunks):

                print(
                    f"Generating embedding "
                    f"{index + 1}/{len(chunks)}..."
                )

                embedding = generate_embedding(
                    chunk
                )

                chunk_records.append({
                    "file_id": file_id,
                    "content": chunk,
                    "embedding": embedding
                })

            # -----------------------------------------
            # SAVE CHUNKS + EMBEDDINGS
            # -----------------------------------------

            supabase \
                .table("document_chunks") \
                .insert(chunk_records) \
                .execute()

            print(
                f"Saved {len(chunks)} "
                f"chunks with embeddings"
            )

        # -----------------------------------------
        # RETURN FILE
        # -----------------------------------------

        return file_record

    except HTTPException:
        raise

    except Exception as error:

        print(
            "Error uploading file:",
            error
        )

        raise HTTPException(
            status_code=503,
            detail="Could not upload file"
        )
@app.get("/api/rooms/{room_id}/documents")
def get_documents(room_id: str):
    room_id = room_id.upper()

    try:
        room_response = (
            supabase
            .table("rooms")
            .select("id")
            .eq("room_code", room_id)
            .execute()
        )

        if not room_response.data:
            raise HTTPException(
                status_code=404,
                detail="Room not found"
            )

        room_db_id = room_response.data[0]["id"]

        response = (
            supabase
            .table("files")
            .select("*")
            .eq("room_id", room_db_id)
            .order("uploaded_at", desc=True)
            .execute()
        )

        return response.data

    except HTTPException:
        raise

    except Exception as error:
        print("Error loading documents:", error)

        raise HTTPException(
            status_code=503,
            detail="Could not load documents"
        )

@app.get("/api/files/{file_id}/download")
def download_file(file_id: str):
    try:
        # Get file metadata
        response = (
            supabase
            .table("files")
            .select("*")
            .eq("id", file_id)
            .execute()
        )

        if not response.data:
            raise HTTPException(
                status_code=404,
                detail="File not found"
            )

        file_info = response.data[0]

        filename = file_info["filename"]
        file_url = file_info["file_url"]

        # Extract storage path from public URL
        storage_path = file_url.split("/project-files/")[-1]

        # Download original file from Supabase Storage
        file_data = (
            supabase
            .storage
            .from_("project-files")
            .download(storage_path)
        )

        return Response(
            content=file_data,
            media_type=file_info["file_type"] or "application/octet-stream",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"'
            }
        )

    except HTTPException:
        raise

    except Exception as error:
        print("Error downloading file:", error)
        raise HTTPException(
            status_code=503,
            detail="Could not download file"
        )

@app.delete("/api/documents/{file_id}")
def delete_document(file_id: str):
    try:
        # Get document metadata
        response = (
            supabase
            .table("files")
            .select("*")
            .eq("id", file_id)
            .execute()
        )

        if not response.data:
            raise HTTPException(
                status_code=404,
                detail="Document not found"
            )

        document = response.data[0]

        filename = document["filename"]
        file_url = document["file_url"]

        # Get storage path from the public URL
        storage_path = file_url.split("/project-files/")[-1]

        # Delete actual file from Supabase Storage
        supabase.storage \
            .from_("project-files") \
            .remove([storage_path])

        # Delete database record
        supabase \
            .table("files") \
            .delete() \
            .eq("id", file_id) \
            .execute()

        return {
            "deleted": True,
            "filename": filename
        }

    except HTTPException:
        raise

    except Exception as error:
        print("Error deleting document:", error)

        raise HTTPException(
            status_code=503,
            detail="Could not delete document"
        )

@app.post("/api/rooms/{room_id}/ask")
def ask_ai(room_id: str, data: dict):

    room_id = room_id.upper()

    question = (data.get("question") or "").strip()

    if not question:
        raise HTTPException(
            status_code=400,
            detail="Question is required"
        )

    try:
        # -----------------------------------------
        # FIND ROOM
        # -----------------------------------------

        room_response = (
            supabase.table("rooms")
            .select("id")
            .eq("room_code", room_id)
            .execute()
        )

        if not room_response.data:
            raise HTTPException(
                status_code=404,
                detail="Room not found"
            )

        room_db_id = room_response.data[0]["id"]

        # -----------------------------------------
        # VECTOR SIMILARITY SEARCH
        # -----------------------------------------

        chunks = search_document_chunks(
            room_db_id,
            question,
            match_count=10
        )

        # -----------------------------------------
        # NO RELEVANT DOCUMENTS
        # -----------------------------------------

        if not chunks:
            return {
                "answer": (
                    "I couldn't find relevant information "
                    "in the documents uploaded to this room."
                ),
                "sources": []
            }

        # -----------------------------------------
        # BUILD CONTEXT + UNIQUE SOURCES
        # -----------------------------------------

        context_parts = []
        source_map = {}

        for index, chunk in enumerate(chunks, start=1):

            file_response = (
                supabase.table("files")
                .select("filename")
                .eq("id", chunk["file_id"])
                .execute()
            )

            filename = "Unknown document"

            if file_response.data:
                filename = file_response.data[0]["filename"]

            # Add chunk to AI context
            context_parts.append(
                f"[Source {index}: {filename}]\n"
                f"{chunk['content']}"
            )

            similarity = round(
                float(chunk["similarity"]),
                3
            )

            # -------------------------------------
            # GROUP SOURCES BY FILE
            # -------------------------------------

            if filename not in source_map:

                source_map[filename] = {
                    "filename": filename,
                    "similarity": similarity,
                    "chunks": 1
                }

            else:

                source_map[filename]["chunks"] += 1

                # Keep highest relevance score
                source_map[filename]["similarity"] = max(
                    source_map[filename]["similarity"],
                    similarity
                )

        # Convert dictionary to list
        sources = list(source_map.values())

        # -----------------------------------------
        # BUILD FINAL CONTEXT
        # -----------------------------------------

        context = "\n\n".join(context_parts)

        # -----------------------------------------
        # RAG PROMPT
        # -----------------------------------------

        prompt = f"""
You are the AI assistant inside a collaborative
document workspace.

Answer the user's question using ONLY the
provided document context.

Rules:
- Use the actual information from the documents.
- Do not simply repeat a document title unless
  the title itself answers the question.
- If the question asks for a problem statement,
  explanation, definition, objective, methodology,
  result, etc., provide the corresponding details
  found in the context.
- Do not invent information.
- Do not use outside knowledge.
- If the answer is not present in the context,
  clearly say that the uploaded documents do not
  contain enough information.
- Mention the relevant document when useful.
- Give a clear and concise answer.

DOCUMENT CONTEXT
================
{context}

USER QUESTION
=============
{question}

ANSWER:
"""

        # -----------------------------------------
        # SEND CONTEXT TO OLLAMA
        # -----------------------------------------

        response = httpx.post(
            "http://127.0.0.1:11434/api/generate",
            json={
                "model": "llama3.2:3b",
                "prompt": prompt,
                "stream": False
            },
            timeout=120
        )

        # -----------------------------------------
        # CHECK OLLAMA RESPONSE
        # -----------------------------------------

        if response.status_code != 200:

            print(
                "Ollama error:",
                response.text
            )

            raise HTTPException(
                status_code=503,
                detail="Local AI model is unavailable"
            )

        result = response.json()

        answer = result.get(
            "response",
            ""
        ).strip()

        if not answer:
            answer = "I couldn't generate an answer."

        # -----------------------------------------
        # RETURN ANSWER + SOURCES
        # -----------------------------------------

        return {
            "answer": answer,
            "sources": sources
        }

    except HTTPException:
        raise

    except Exception as error:

        print(
            "RAG ERROR:",
            error
        )

        raise HTTPException(
            status_code=503,
            detail="Could not process the AI request"
        )